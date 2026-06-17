"""Excel and PDF exporters for the JDT Pricing Model.

The front-end posts already-computed rows (it owns the calculation engine), so
these builders only format and lay out the supplied values — Excel/PDF output
always matches exactly what the user sees on screen.
"""
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.enums import TA_RIGHT, TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .pdf_quote import safe_filename

NAVY = "26395B"
NAVY2 = "34507A"
ICE = "EAF4FA"
_NAVY = colors.HexColor("#26395b")
_ICE = colors.HexColor("#eaf4fa")
_LINE = colors.HexColor("#d8e2ec")
_MUTED = colors.HexColor("#6b7686")

_HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
_HDR_FILL = PatternFill("solid", fgColor=NAVY)
_TITLE_FONT = Font(color="26395B", bold=True, size=14)
_BOLD = Font(bold=True)
_THIN = Side(style="thin", color="E5E9F0")
_BORDER = Border(bottom=_THIN)


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------
def _header_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        c.font = _HDR_FONT
        c.fill = _HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")


def _autofit(ws, max_width=46):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 0
        for cell in col:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(11, longest + 2), max_width)


MONEY = '#,##0.00'
MONEY0 = '#,##0'
PCT = '0.0%'


def _set(ws, row, col, value, fmt=None, bold=False, align=None):
    c = ws.cell(row=row, column=col, value=value)
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = _BOLD
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


# ---------------------------------------------------------------------------
# Excel: full model
# ---------------------------------------------------------------------------
def export_full_excel(payload, out_path):
    wb = Workbook()

    # --- Settings ---
    ws = wb.active
    ws.title = "Settings"
    ws["A1"] = "JDT Pricing Model — Settings"
    ws["A1"].font = _TITLE_FONT
    s = payload.get("settings", {})
    rows = [
        ("Driver base hourly ($/hr)", s.get("driverBaseHourly"), MONEY),
        ("OT multiplier", s.get("otMultiplier"), '0.0000'),
        ("Public holiday hourly ($/hr)", s.get("publicHolidayHourly"), MONEY),
        ("Linehaul hourly ($/hr)", s.get("linehaulHourly"), MONEY),
        ("Superannuation", s.get("superRate"), '0.0000'),
        ("Workcover", s.get("workcoverRate"), '0.0000'),
        ("Payroll tax", s.get("payrollTaxRate"), '0.0000'),
        ("Vehicle $/km — Ute", s.get("vehicleRateUte"), MONEY),
        ("Vehicle $/km — Rigid", s.get("vehicleRateRigid"), MONEY),
        ("Vehicle $/km — Semi", s.get("vehicleRateSemi"), MONEY),
        ("Vehicle $/km — Bdouble", s.get("vehicleRateBdouble"), MONEY),
        ("Fuel levy %", s.get("fuelLevyPct"), '0.0'),
        ("Target margin %", s.get("targetMarginPct"), '0.0'),
    ]
    r = 3
    for label, val, fmt in rows:
        ws.cell(row=r, column=1, value=label).font = _BOLD
        _set(ws, r, 2, val, fmt)
        r += 1
    _autofit(ws)

    # --- Lanes ---
    ws = wb.create_sheet("Lanes")
    headers = ["Origin", "Destination", "Vehicle", "Spaces", "Trips/wk", "Hrs/trip", "Km/trip",
               "Tolls $", "Overnight $", "Load extras $", "Cost/trip $", "Base price $",
               "Price+FL $", "Margin %", "Annual rev $", "Annual GP $", "Decision", "Quote?"]
    _header_row(ws, 1, headers)
    r = 2
    for L in payload.get("lanes", []):
        vals = [L.get("origin"), L.get("dest"), L.get("vehicle"), L.get("spaces"), L.get("trips"),
                L.get("hrs"), L.get("km"), L.get("tolls"), L.get("overnight"), L.get("loadExtras"),
                L.get("cost"), L.get("base"), L.get("priceFL"), L.get("margin"),
                L.get("annualRev"), L.get("annualGP"), L.get("decision"), L.get("quote")]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=i + 1, value=v)
            if 10 <= i <= 12:
                c.number_format = MONEY
            elif i == 13:
                c.number_format = PCT
            elif i in (14, 15):
                c.number_format = MONEY0
        r += 1
    ws.freeze_panes = "A2"
    _autofit(ws)

    # --- Warehousing ---
    ws = wb.create_sheet("Warehousing")
    headers = ["Customer", "Pallets", "Inb/wk", "Outb/wk", "Cases/wk", "VAS hrs/wk", "Other $/wk",
               "Cost/wk $", "Base price $/wk", "Margin %", "$/pallet", "Annual rev $", "Annual GP $",
               "Decision", "Quote?"]
    _header_row(ws, 1, headers)
    r = 2
    for A in payload.get("accounts", []):
        vals = [A.get("customer"), A.get("pallets"), A.get("inb"), A.get("outb"), A.get("cases"),
                A.get("vasHrs"), A.get("other"), A.get("cost"), A.get("base"), A.get("margin"),
                A.get("perPallet"), A.get("annualRev"), A.get("annualGP"), A.get("decision"), A.get("quote")]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=i + 1, value=v)
            if i in (7, 8, 10):
                c.number_format = MONEY
            elif i == 9:
                c.number_format = PCT
            elif i in (11, 12):
                c.number_format = MONEY0
        r += 1
    ws.freeze_panes = "A2"
    _autofit(ws)

    # --- Summary ---
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Portfolio Summary"
    ws["A1"].font = _TITLE_FONT
    sm = payload.get("summary", {})

    def block(title, d, row):
        ws.cell(row=row, column=1, value=title).font = _BOLD
        items = [
            ("Weekly revenue", d.get("weekRev"), MONEY0),
            ("Weekly cost", d.get("weekCost"), MONEY0),
            ("Weekly GP", d.get("weekGP"), MONEY0),
            ("Blended margin %", d.get("margin"), PCT),
            ("Annual revenue", d.get("annualRev"), MONEY0),
            ("Annual GP", d.get("annualGP"), MONEY0),
        ]
        rr = row + 1
        for label, val, fmt in items:
            ws.cell(row=rr, column=1, value=label)
            _set(ws, rr, 2, val, fmt)
            rr += 1
        return rr + 1

    nxt = block("TRANSPORT", sm.get("transport", {}), 3)
    nxt = block("WAREHOUSING", sm.get("warehousing", {}), nxt)
    block("COMBINED", sm.get("combined", {}), nxt)
    _autofit(ws)

    # --- Tenders ---
    _tenders_sheet(wb, payload)

    wb.save(out_path)


def _tenders_sheet(wb, payload, title="Tenders"):
    ws = wb.create_sheet(title)
    headers = ["Reference", "Customer", "Title", "Status", "Due date", "Submitted",
               "Value $", "Probability %", "Weighted $", "Owner", "Saved?", "Updated", "Notes"]
    _header_row(ws, 1, headers)
    r = 2
    for t in payload.get("tenders", []):
        try:
            value = float(t.get("value") or 0)
        except (TypeError, ValueError):
            value = 0
        try:
            prob = float(t.get("probability") or 0)
        except (TypeError, ValueError):
            prob = 0
        vals = [t.get("reference"), t.get("customer"), t.get("title"), t.get("status"),
                t.get("dueDate"), t.get("submittedDate"), value, prob / 100.0, value * prob / 100.0,
                t.get("owner"), "Yes" if t.get("snapshot") else "No", t.get("updatedAt"), t.get("notes")]
        for i, v in enumerate(vals):
            c = ws.cell(row=r, column=i + 1, value=v)
            if i in (6, 8):
                c.number_format = MONEY0
            elif i == 7:
                c.number_format = PCT
        r += 1
    # pipeline totals
    p = payload.get("pipeline", {})
    if p:
        r += 1
        ws.cell(row=r, column=1, value="PIPELINE").font = _BOLD
        summary = [
            ("Open tenders", p.get("openCount")), ("Open value $", p.get("openValue")),
            ("Weighted value $", p.get("weightedValue")), ("Won value $", p.get("wonValue")),
            ("Lost value $", p.get("lostValue")),
            ("Win rate %", p.get("winRate")), ("Due ≤14 days", p.get("dueSoon")),
            ("Overdue", p.get("overdue")),
        ]
        for label, val in summary:
            r += 1
            ws.cell(row=r, column=1, value=label).font = _BOLD
            c = ws.cell(row=r, column=2, value=val)
            if "$" in label:
                c.number_format = MONEY0
            elif "%" in label:
                c.number_format = PCT
    ws.freeze_panes = "A2"
    _autofit(ws)
    return ws


def export_tenders_excel(payload, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    _tenders_sheet(wb, payload, title="Tender Register")
    wb.save(out_path)


# Column headers match the import auto-mapper keywords so an uploaded copy of
# this template maps onto JDT lane fields with no manual mapping.
TEMPLATE_HEADERS = [
    "Collection postcode", "Collection suburb", "Delivery postcode", "Delivery suburb",
    "Pallets", "Weight (kg)", "Dimensions (LxWxH)", "Stackable (Y/N)", "Loading type",
    "Vehicle", "Frequency per week", "Hours per trip", "Km per trip",
]
TEMPLATE_EXAMPLES = [
    ["4110", "Acacia Ridge", "4116", "Sunnybank", 6, 420, "1.2x1.0x1.6", "Y", "Tail-lift", "Rigid", 5, 4, 60],
    ["4110", "Acacia Ridge", "4350", "Toowoomba", 22, 9000, "Std pallet", "Y", "Forklift", "Semi", 3, 9, 130],
]


def export_import_template(payload, out_path):
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Tender Lanes"
    _header_row(ws, 1, TEMPLATE_HEADERS)
    ws.freeze_panes = "A2"
    for r, row in enumerate(TEMPLATE_EXAMPLES, start=2):
        for cidx, val in enumerate(row, start=1):
            ws.cell(row=r, column=cidx, value=val)

    # dropdowns for the constrained columns (rows 2..500)
    def add_list(col_letter, options):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("%s2:%s500" % (col_letter, col_letter))

    add_list("H", ["Y", "N"])
    add_list("I", ["Tail-lift", "Dock / Ramp", "Forklift", "Hand unload", "Crane", "Side-loader"])
    add_list("J", ["Ute", "Rigid", "Semi", "Bdouble"])
    _autofit(ws)

    # instructions sheet
    info = wb.create_sheet("How to use")
    info["A1"] = "JDT Tender Lane Import Template"
    info["A1"].font = _TITLE_FONT
    notes = [
        "",
        "1. Enter one row per lane on the 'Tender Lanes' sheet (example rows included — overwrite or delete them).",
        "2. Keep the header row as-is so the importer maps the columns automatically.",
        "3. Postcodes/suburbs are used to plot the lane network on the map (needs internet).",
        "4. Pricing fields the cost engine uses: Pallets, Vehicle, Frequency per week, Hours per trip, Km per trip.",
        "5. Operational-only fields (Weight, Dimensions, Stackable, Loading type) are stored for reference.",
        "6. Save the file, then in the app open a tender -> Operational Lanes -> Import from Excel -> upload this file.",
        "",
        "Vehicle options:  Ute, Rigid, Semi, Bdouble",
        "Loading options:  Tail-lift, Dock / Ramp, Forklift, Hand unload, Crane, Side-loader",
    ]
    for i, line in enumerate(notes, start=2):
        info.cell(row=i, column=1, value=line)
    info.column_dimensions["A"].width = 100
    wb.save(out_path)


def export_quote_excel(payload, out_path):
    quote = payload.get("quote", {})
    comp = payload.get("computed", {})
    wb = Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws["A1"] = "JD Refrigerated Transport — Quotation"
    ws["A1"].font = _TITLE_FONT
    meta = [
        ("Customer", quote.get("customer")), ("Contact", quote.get("contact")),
        ("Reference", quote.get("reference")), ("Quote #", quote.get("quoteNumber")),
        ("Issued", datetime.date.today().strftime("%d %b %Y")),
        ("Payment terms", quote.get("terms") or "14 days from invoice"),
    ]
    r = 3
    for k, v in meta:
        ws.cell(row=r, column=1, value=k).font = _BOLD
        ws.cell(row=r, column=2, value=v)
        r += 1

    r += 1
    if comp.get("transport"):
        ws.cell(row=r, column=1, value="TRANSPORT — LINEHAUL & METRO LANES").font = _BOLD
        r += 1
        _header_row(ws, r, ["Origin", "Destination", "Vehicle", "Spaces", "Trips/wk", "Rate/trip", "Weekly $", "Annual $"])
        r += 1
        for x in comp["transport"]:
            cells = [x.get("origin"), x.get("dest"), x.get("vehicle"), x.get("spaces"), x.get("trips"),
                     x.get("ratePerTrip"), x.get("weekly"), x.get("annual")]
            for i, v in enumerate(cells):
                c = ws.cell(row=r, column=i + 1, value=v)
                if i in (5, 6):
                    c.number_format = MONEY
                elif i == 7:
                    c.number_format = MONEY0
            r += 1
        ws.cell(row=r, column=1, value="Transport subtotal").font = _BOLD
        _set(ws, r, 7, comp.get("transportWeekly"), MONEY, bold=True)
        _set(ws, r, 8, comp.get("transportAnnual"), MONEY0, bold=True)
        r += 2

    if comp.get("warehousing"):
        ws.cell(row=r, column=1, value="WAREHOUSING — STORAGE, HANDLING & VAS").font = _BOLD
        r += 1
        _header_row(ws, r, ["Customer", "Pallets", "Inb/wk", "Outb/wk", "Cases/wk", "$/pallet", "Weekly $", "Annual $"])
        r += 1
        for x in comp["warehousing"]:
            cells = [x.get("customer"), x.get("pallets"), x.get("inb"), x.get("outb"), x.get("cases"),
                     x.get("perPallet"), x.get("weekly"), x.get("annual")]
            for i, v in enumerate(cells):
                c = ws.cell(row=r, column=i + 1, value=v)
                if i in (5, 6):
                    c.number_format = MONEY
                elif i == 7:
                    c.number_format = MONEY0
            r += 1
        ws.cell(row=r, column=1, value="Warehousing subtotal").font = _BOLD
        _set(ws, r, 7, comp.get("warehousingWeekly"), MONEY, bold=True)
        _set(ws, r, 8, comp.get("warehousingAnnual"), MONEY0, bold=True)
        r += 2

    ws.cell(row=r, column=1, value="TOTAL CONTRACT VALUE").font = _TITLE_FONT
    _set(ws, r, 7, comp.get("totalWeekly"), MONEY, bold=True)
    _set(ws, r, 8, comp.get("totalAnnual"), MONEY0, bold=True)
    _autofit(ws)
    wb.save(out_path)


# ---------------------------------------------------------------------------
# PDF helpers
# ---------------------------------------------------------------------------
def _pdf_styles():
    ss = getSampleStyleSheet()
    return {
        "h1": ParagraphStyle("h1", parent=ss["Title"], fontSize=18, textColor=_NAVY, spaceAfter=2),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12, textColor=_NAVY, spaceBefore=10, spaceAfter=4),
        "small": ParagraphStyle("s", parent=ss["Normal"], fontSize=8, textColor=_MUTED),
    }


def _money(v, dp=0):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return ""
    return ("-$" if v < 0 else "$") + f"{abs(v):,.{dp}f}"


def _pct(v):
    try:
        return f"{float(v) * 100:.1f}%"
    except (TypeError, ValueError):
        return ""


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return "" if v is None else str(v)
    return f"{int(f):,}" if f == int(f) else f"{f:,.2f}"


def _grid_table(headers, rows, col_widths, align_from):
    data = [headers] + rows
    t = Table(data, repeatRows=1, colWidths=[w * mm for w in col_widths], hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ALIGN", (align_from, 0), (-1, -1), "RIGHT"),
        ("ALIGN", (0, 0), (align_from - 1, -1), "LEFT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, _ICE]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.3, _LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _doc(out_path, title, landscape_mode=True):
    size = landscape(A4) if landscape_mode else A4
    return SimpleDocTemplate(out_path, pagesize=size, leftMargin=12 * mm, rightMargin=12 * mm,
                             topMargin=12 * mm, bottomMargin=12 * mm, title=title)


def export_data_pdf(payload, out_path):
    st = _pdf_styles()
    story = [Paragraph("JDT Pricing Model — Working Data", st["h1"]),
             Paragraph("Generated " + datetime.date.today().strftime("%d %b %Y"), st["small"])]

    lanes = payload.get("lanes", [])
    if lanes:
        story.append(Paragraph("TRANSPORT LANES", st["h2"]))
        rows = [[L.get("origin", ""), L.get("dest", ""), L.get("vehicle", ""), _num(L.get("spaces")),
                 _num(L.get("trips")), _num(L.get("hrs")), _num(L.get("km")), _money(L.get("cost"), 2),
                 _money(L.get("base"), 2), _money(L.get("priceFL"), 2), _pct(L.get("margin")),
                 _money(L.get("annualRev")), _money(L.get("annualGP")), L.get("decision", "")] for L in lanes]
        story.append(_grid_table(
            ["Origin", "Destination", "Veh", "Sp", "Trips", "Hrs", "Km", "Cost", "Base", "Price+FL", "Marg", "Ann rev", "Ann GP", "Dec"],
            rows, [30, 32, 14, 11, 13, 12, 14, 20, 20, 22, 16, 24, 24, 20], align_from=3))

    accts = payload.get("accounts", [])
    if accts:
        story.append(Paragraph("WAREHOUSING ACCOUNTS", st["h2"]))
        rows = [[A.get("customer", ""), _num(A.get("pallets")), _num(A.get("inb")), _num(A.get("outb")),
                 _num(A.get("cases")), _num(A.get("vasHrs")), _money(A.get("cost"), 2), _money(A.get("base"), 2),
                 _pct(A.get("margin")), _money(A.get("perPallet"), 2), _money(A.get("annualRev")),
                 _money(A.get("annualGP")), A.get("decision", "")] for A in accts]
        story.append(_grid_table(
            ["Customer", "Pallets", "Inb", "Outb", "Cases", "VAS h", "Cost/wk", "Base/wk", "Marg", "$/pallet", "Ann rev", "Ann GP", "Dec"],
            rows, [40, 18, 14, 16, 16, 16, 22, 22, 16, 22, 26, 26, 18], align_from=1))

    if not lanes and not accts:
        story.append(Paragraph("No priced lanes or accounts to export.", st["small"]))

    _doc(out_path, "JDT Working Data").build(story)


def export_tenders_pdf(payload, out_path):
    st = _pdf_styles()
    story = [Paragraph("JDT Tender Register", st["h1"]),
             Paragraph("Generated " + datetime.date.today().strftime("%d %b %Y"), st["small"]), Spacer(1, 6)]

    p = payload.get("pipeline", {})
    if p:
        cards = [
            ["Open", str(p.get("openCount", 0))], ["Open value", _money(p.get("openValue"))],
            ["Weighted", _money(p.get("weightedValue"))], ["Won", _money(p.get("wonValue"))],
            ["Win rate", _pct(p.get("winRate"))], ["Due ≤14d", str(p.get("dueSoon", 0))],
            ["Overdue", str(p.get("overdue", 0))],
        ]
        ct = Table([[c[0] for c in cards], [c[1] for c in cards]], colWidths=[38 * mm] * len(cards))
        ct.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), _NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("BACKGROUND", (0, 1), (-1, 1), _ICE), ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(ct)
        story.append(Spacer(1, 10))

    rows = []
    for t in payload.get("tenders", []):
        try:
            value = float(t.get("value") or 0)
            prob = float(t.get("probability") or 0)
        except (TypeError, ValueError):
            value, prob = 0, 0
        rows.append([t.get("reference", ""), t.get("customer", ""), t.get("title", ""), t.get("status", ""),
                     t.get("dueDate", ""), _money(value), _pct(prob / 100.0), _money(value * prob / 100.0),
                     t.get("owner", "")])
    if rows:
        story.append(_grid_table(
            ["Reference", "Customer", "Title", "Status", "Due", "Value", "Prob", "Weighted", "Owner"],
            rows, [34, 38, 40, 24, 24, 28, 18, 28, 28], align_from=5))
    else:
        story.append(Paragraph("No tenders to export.", st["small"]))

    _doc(out_path, "JDT Tender Register").build(story)


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------
_DISPATCH = {
    "excel-full": (export_full_excel, "JDT_Model", "xlsx"),
    "data-pdf": (export_data_pdf, "JDT_Data", "pdf"),
    "quote-excel": (export_quote_excel, "JDT_Quote", "xlsx"),
    "tenders-excel": (export_tenders_excel, "JDT_Tenders", "xlsx"),
    "tenders-pdf": (export_tenders_pdf, "JDT_Tenders", "pdf"),
    "import-template": (export_import_template, "JDT_Tender_Import_Template", "xlsx"),
}


def run_export(payload, dest_dir):
    """Build the requested export and return (filename, full_path)."""
    import os
    etype = payload.get("type")
    if etype not in _DISPATCH:
        raise ValueError(f"unknown export type: {etype}")
    builder, stem, ext = _DISPATCH[etype]
    # tag quote exports with the customer for easy filing
    if etype == "quote-excel":
        cust = safe_filename((payload.get("quote") or {}).get("customer") or "Quote")
        stem = f"JDT_Quote_{cust}"
    date = datetime.date.today().strftime("%Y%m%d")
    filename = f"{stem}_{date}.{ext}"
    full = os.path.join(dest_dir, filename)
    builder(payload, full)
    return filename, full
